"""Generate 5 equation PNGs plus a sample.docx with them and one non-equation image."""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
matplotlib.rcParams["mathtext.fontset"] = "cm"  # LaTeX look, like KaTeX/MathJax output
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent
EQS = [
    r"E = mc^2",
    r"x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}",
    r"\int_0^\infty e^{-x^2}\,dx = \frac{\sqrt{\pi}}{2}",
    r"\sum_{n=1}^{\infty} \frac{1}{n^2} = \frac{\pi^2}{6}",
    r"f(x) = \alpha x^3 + \beta \sin(\theta)",
]

paths = []
for i, eq in enumerate(EQS):
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${eq}$", fontsize=20)
    p = HERE / f"eq{i}.png"
    fig.savefig(p, dpi=200, bbox_inches="tight", pad_inches=0.08, facecolor="white")
    plt.close(fig)
    paths.append(p)

photo = HERE / "photo.png"
plt.imsave(photo, np.random.default_rng(0).random((120, 160, 3)))

doc = Document()
for i, p in enumerate(paths):
    para = doc.add_paragraph(f"Equation {i + 1}: ")
    para.add_run().add_picture(str(p), height=Pt(22))
doc.add_paragraph("A non-equation picture:").add_run().add_picture(str(photo), height=Pt(60))
doc.save(HERE / "sample.docx")
print("ok")
