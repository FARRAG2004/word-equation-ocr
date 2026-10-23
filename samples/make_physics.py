"""Physics / electrical benchmark with exact ground truth: every image is rendered from its own LaTeX.

python samples/make_physics.py  ->  samples/physics_bench.docx + samples/ground_truth_physics.json
Each equation is rendered in three fonts (Computer Modern like KaTeX/MathJax, STIX like Times,
DejaVu Sans) and two resolutions, as AI-generated documents vary.
"""
import io
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent
EQS = {
    "electrical": [
        r"U=IR", r"P=UI\cos\varphi", r"Q=UI\sin\varphi", r"S=\sqrt{P^2+Q^2}",
        r"X_L=\omega L=2\pi fL", r"X_C=\frac{1}{\omega C}", r"Z=\sqrt{R^2+(X_L-X_C)^2}",
        r"\varphi=\arctan\frac{X_L-X_C}{R}", r"R_{eq}=R_1+R_2+R_3", r"\frac{1}{R_{eq}}=\frac{1}{R_1}+\frac{1}{R_2}",
        r"\sum_{k=1}^{n}I_k=0", r"\sum_{k=1}^{n}U_k=0", r"i(t)=I_m\sin(\omega t+\psi)",
        r"U_{rms}=\frac{U_m}{\sqrt{2}}", r"\tau=RC", r"u_C(t)=U_0\left(1-e^{-t/\tau}\right)",
        r"\eta=\frac{P_2}{P_1}\cdot 100\%", r"E=c_e\Phi n", r"M=c_M\Phi I_a",
        r"\omega=\frac{U-I_aR_a}{c_e\Phi}", r"I_a=\frac{U-E}{R_a+R_{ad}}", r"n_1=\frac{60f_1}{p}",
        r"s=\frac{n_1-n_2}{n_1}", r"k_{tr}=\frac{w_1}{w_2}=\frac{E_1}{E_2}", r"P_{Cu}=3I_1^2R_1",
        r"W=\int_0^t u(t)i(t)\,dt", r"\underline{Z}=R+jX", r"f_0=\frac{1}{2\pi\sqrt{LC}}",
        r"Q=\frac{1}{R}\sqrt{\frac{L}{C}}", r"\Phi=\frac{\mu_0 N I A}{l}",
    ],
    "physics": [
        r"F=ma", r"E=mc^2", r"p=mv", r"E_k=\frac{1}{2}mv^2", r"F=G\frac{m_1m_2}{r^2}",
        r"F=k\frac{q_1q_2}{r^2}", r"\vec{F}=q\vec{E}+q\vec{v}\times\vec{B}", r"\nabla\cdot\mathbf{E}=\frac{\rho}{\varepsilon_0}",
        r"\nabla\times\mathbf{E}=-\frac{\partial\mathbf{B}}{\partial t}", r"\oint_C\mathbf{B}\cdot d\mathbf{l}=\mu_0 I",
        r"\lambda=\frac{h}{p}", r"E=h\nu", r"\Delta x\,\Delta p\geq\frac{\hbar}{2}",
        r"PV=nRT", r"\Delta U=Q-W", r"\eta=1-\frac{T_c}{T_h}", r"v=\lambda f",
        r"T=2\pi\sqrt{\frac{l}{g}}", r"x(t)=A\cos(\omega t+\varphi_0)", r"\omega=\sqrt{\frac{k}{m}}",
        r"\frac{d^2x}{dt^2}+2\beta\frac{dx}{dt}+\omega_0^2x=0", r"I=I_0e^{-\mu x}", r"N=N_0e^{-\lambda t}",
        r"\rho=\frac{m}{V}", r"P=\frac{F}{A}", r"\gamma=\frac{1}{\sqrt{1-v^2/c^2}}", r"\Phi_B=\int\mathbf{B}\cdot d\mathbf{A}",
        r"\varepsilon=-\frac{d\Phi_B}{dt}", r"\mathbf{J}=\sigma\mathbf{E}", r"\dot{q}=-k\nabla T",
    ],
}
FONTS = {"cm": "cm", "stix": "stix", "dejavu": "dejavusans"}
DPIS = (300, 150)


def render(latex, fontset, dpi):
    with matplotlib.rc_context({"mathtext.fontset": fontset}):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=16)
        buf = io.BytesIO()
        fig.savefig(buf, dpi=dpi, bbox_inches="tight", pad_inches=0.06, facecolor="white")
        plt.close(fig)
    return buf


doc, truth, n = Document(), {}, 0
doc.add_heading("Physics / electrical benchmark", 1)
for cat, eqs in EQS.items():
    for latex in eqs:
        for font, fontset in FONTS.items():
            for dpi in DPIS:
                n += 1
                eq_id = f"EQ-{n:03d}"
                doc.add_paragraph(f"{eq_id}  [{cat}-{font}-{dpi}]")
                doc.add_paragraph().add_run().add_picture(render(latex, fontset, dpi), height=Pt(20))
                truth[eq_id] = latex
doc.save(HERE / "physics_bench.docx")
(HERE / "ground_truth_physics.json").write_text(json.dumps(truth, indent=1), encoding="utf-8")
print(f"{n} images")
