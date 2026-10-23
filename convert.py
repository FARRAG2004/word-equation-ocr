"""CLI fallback: python convert.py input.docx output.docx [--no-verify]"""
import argparse
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).parent / "server"))
import core  # noqa: E402


def paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--no-verify", action="store_true",
                    help="convert every readable image, to measure raw OCR accuracy")
    args = ap.parse_args()

    doc = Document(args.src)
    counts = Counter({"converted": 0, "flagged": 0, "skipped": 0})
    would_skip = 0
    for i, p in enumerate(paragraphs(doc)):
        for run in list(p._p.iter(qn("w:r"))):
            blip = run.find(f".//{qn('wp:inline')}//{qn('a:blip')}")
            if blip is None:
                continue
            part = doc.part.related_parts[blip.get(qn("r:embed"))]
            r = core.process(part.blob, verify=not args.no_verify, source=f"{Path(args.src).name}:{part.partname}")
            counts[r["status"]] += 1
            if r["status"] != "converted":
                print(f"{r['status'].upper()} {part.partname} ({r['reason']}): {r['latex']}")
            if r["status"] == "converted":
                run.addprevious(etree.fromstring(r["omml"]))
                run.getparent().remove(run)
    doc.save(args.dst)
    print(dict(counts))
    print(f"per-image log: {Path(core.LOG_PATH).resolve()}")


if __name__ == "__main__":
    main()
